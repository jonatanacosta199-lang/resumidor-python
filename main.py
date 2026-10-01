import os
import uuid
import fitz  # PyMuPDF
import docx
import google.generativeai as genai
from fastapi import FastAPI, UploadFile, File, Form, BackgroundTasks, HTTPException
from fastapi.responses import HTMLResponse, FileResponse
from fastapi.templating import Jinja2Templates
from fastapi.requests import Request

app = FastAPI(title="Resumidor Académico Pro")
templates = Jinja2Templates(directory="templates")

JOBS = {}

def process_pdf_job(job_id: str, file_bytes: bytes, filename: str, api_key: str, ratio: float):
    try:
        JOBS[job_id]["status"] = "processing"
        JOBS[job_id]["message"] = "Extrayendo texto del documento..."

        genai.configure(api_key=api_key)
        model = genai.GenerativeModel("gemini-1.5-flash")

        doc = fitz.open(stream=file_bytes, filetype="pdf")
        total_pages = len(doc)
        extracted_chunks = []

        for page_num in range(total_pages):
            page = doc[page_num]
            text = page.get_text("text").strip()

            # Si es fotocopia o escaneo sin texto digital, usar la visión multimodal de Gemini
            if len(text) < 50:
                pix = page.get_pixmap(dpi=150)
                img_bytes = pix.tobytes("png")
                ocr_prompt = "Extrae de forma limpia todo el contenido visible en esta página escaneada, omitiendo encabezados o números de página."
                response = model.generate_content([ocr_prompt, {"mime_type": "image/png", "data": img_bytes}])
                text = response.text if response.text else ""

            if text:
                extracted_chunks.append(f"--- Página {page_num + 1} ---\n{text}")

            progress = int(((page_num + 1) / total_pages) * 50)
            JOBS[job_id]["progress"] = progress
            JOBS[job_id]["message"] = f"Leyendo página {page_num + 1} de {total_pages}..."

        full_text = "\n\n".join(extracted_chunks)

        if not full_text.strip():
            raise Exception("No se pudo extraer texto del documento.")

        JOBS[job_id]["message"] = "Generando resumen académico..."
        JOBS[job_id]["progress"] = 60

        pct = int(ratio * 100)
        system_instruction = f"""Eres un historiador y sintetizador académico.
Genera un resumen estructurado, riguroso e hiperfiel al texto.
- Utiliza títulos '### ' para organizar los temas o capítulos.
- Conserva aproximadamente el {pct}% de los conceptos clave, nombres, fechas e hipótesis.
- Comienza directamente con el contenido sin introducciones vacías."""

        prompt = f"Título del Documento: {filename}\n\nContenido:\n{full_text}"
        
        final_response = model.generate_content(f"{system_instruction}\n\n{prompt}")
        summary_text = final_response.text

        # Generar Word (.docx)
        docx_doc = docx.Document()
        docx_doc.add_heading(filename, level=1)
        for paragraph in summary_text.split("\n\n"):
            p_text = paragraph.strip()
            if p_text.startswith("###"):
                docx_doc.add_heading(p_text.replace("###", "").strip(), level=2)
            elif p_text:
                docx_doc.add_paragraph(p_text)

        os.makedirs("downloads", exist_ok=True)
        docx_filename = f"downloads/{job_id}.docx"
        docx_doc.save(docx_filename)

        JOBS[job_id]["status"] = "completed"
        JOBS[job_id]["progress"] = 100
        JOBS[job_id]["message"] = "¡Resumen finalizado!"
        JOBS[job_id]["result"] = summary_text
        JOBS[job_id]["docx_path"] = docx_filename

    except Exception as e:
        JOBS[job_id]["status"] = "failed"
        JOBS[job_id]["error"] = str(e)

@app.get("/", response_class=HTMLResponse)
async def index(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})

@app.post("/api/summarize")
async def start_summarize(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    api_key: str = Form(...),
    ratio: float = Form(0.70)
):
    if not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=400, detail="Solo se admiten archivos PDF.")

    file_bytes = await file.read()
    job_id = str(uuid.uuid4())
    
    JOBS[job_id] = {
        "status": "queued",
        "progress": 0,
        "message": "En cola...",
        "filename": file.filename
    }

    background_tasks.add_task(process_pdf_job, job_id, file_bytes, file.filename, api_key, ratio)
    return {"job_id": job_id}

@app.get("/api/status/{job_id}")
async def get_status(job_id: str):
    if job_id not in JOBS:
        raise HTTPException(status_code=404, detail="Tarea no encontrada")
    return JOBS[job_id]

@app.get("/api/download/{job_id}")
async def download_file(job_id: str):
    job = JOBS.get(job_id)
    if not job or job.get("status") != "completed":
        raise HTTPException(status_code=404, detail="Archivo no disponible")
    return FileResponse(
        job["docx_path"],
        filename=f"Resumen_{job['filename']}.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )